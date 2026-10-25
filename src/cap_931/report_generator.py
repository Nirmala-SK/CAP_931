import io
import re
from datetime import date
from pathlib import Path
from urllib.parse import urlparse

from docx import Document
from docx.enum.text import WD_UNDERLINE
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

TEMPLATE = Path(__file__).parent / "perscholas_template.docx"
FONT = "Arial"
BODY_PT = 8.5


def build_links_section(pages, news, filing=None, competitor_pages=None):
    """Build the Article Links section from the sources we actually used."""
    lines = ["## Article Links", "", "**Company pages**"]
    for url in pages:
        lines.append(f"- [{url}]({url})")

    if filing:
        label = f"{filing['company']} {filing['form']} filed {filing['date']}"
        lines += ["", "**SEC filing**", f"- [{label}]({filing['url']})"]

    if competitor_pages:
        lines += ["", "**Competitor pages**"]
        for name, comp_pages in competitor_pages.items():
            for url in comp_pages:
                path = urlparse(url).path or "/"
                lines.append(f"- [{name}: {path}]({url})")

    if news:
        lines += ["", "**Recent news**"]
        for a in news:
            title = a["title"].replace("[", "(").replace("]", ")")
            suffix = f" - {a['source']}"
            if title.endswith(suffix):
                title = title[: -len(suffix)]
            date_text = a["date"][:16]
            lines.append(f"- [{title}]({a['link']}) ({a['source']}, {date_text})")

    return "\n".join(lines)


def _style_run(run, size=BODY_PT, bold=False, italic=False):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def _add_hyperlink(paragraph, text, url, size=BODY_PT):
    """Add a clickable link to a paragraph."""
    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = paragraph.add_run(text)
    _style_run(run, size)
    run.font.color.rgb = RGBColor(0x03, 0x9B, 0xE5)
    run.font.underline = WD_UNDERLINE.SINGLE
    link.append(run._element)
    paragraph._p.append(link)


TOKEN = re.compile(r"(\*\*.+?\*\*|\[[^\]]+?\]\([^)]+?\)|(?<!\*)\*[^*\s][^*]*?\*(?!\*))")


def _add_inline(paragraph, text, size=BODY_PT):
    """Write text into a paragraph, handling **bold**, *italic* and [text](url)."""
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            _style_run(paragraph.add_run(part[2:-2]), size, bold=True)
        elif part.startswith("[") and "](" in part:
            label, url = re.match(r"\[(.+?)\]\((.+?)\)$", part).groups()
            _add_hyperlink(paragraph, label, url, size)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            _style_run(paragraph.add_run(part[1:-1]), size, italic=True)
        else:
            _style_run(paragraph.add_run(part), size)


def _spacing(paragraph, before=0, after=2):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.0


MAX_NEWS_IN_DOCX = 5
SECTION_CAPS = {"**Recent news**": MAX_NEWS_IN_DOCX}
# These link lists are written on one line each in the Word file to save space.
INLINE_SECTIONS = {"**Company pages**": "Company pages", "**Competitor pages**": "Competitor pages"}


def _short_label(url, name=""):
    """Turn a web address into a short readable label."""
    path = urlparse(url).path.strip("/")
    last = path.split("/")[-1] if path else ""
    text = last.replace("-", " ").replace("_", " ").title() if last else "Home"
    if len(text) > 26:
        text = text[:25].rstrip() + "\u2026"
    return f"{name}: {text}" if name else text


def _flush_links(doc, title, links):
    """Write collected links as one paragraph: Title:  link | link | link."""
    if not title or not links:
        return
    p = doc.add_paragraph()
    _style_run(p.add_run(f"{title}:  "), BODY_PT, bold=True)
    for k, (label, url) in enumerate(links):
        if k:
            _style_run(p.add_run("  |  "), BODY_PT)
        _add_hyperlink(p, label, url)
    _spacing(p, 0, 2)


def _remove_template_notes(doc):
    """Blank out any 'TEMPLATE ... make a copy' note left in the template headers."""
    for section in doc.sections:
        for header in (section.header, section.first_page_header, section.even_page_header):
            for text in header._element.iter(qn("w:t")):
                if text.text and ("TEMPLATE" in text.text or "make a copy" in text.text):
                    text.text = ""


def build_docx(brief_markdown, company_name, product_name, author=""):
    """Turn the brief into a one-page Word document using the PerScholas template."""
    doc = Document(str(TEMPLATE))
    _remove_template_notes(doc)

    body = doc.element.body
    for element in list(body):
        if not element.tag.endswith("sectPr"):
            body.remove(element)

    section = doc.sections[0]
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    title = doc.add_paragraph(style="Title")
    _style_run(title.add_run(f"Account Brief: {company_name}"), 20)
    _spacing(title, 0, 2)

    sub = doc.add_paragraph()
    byline = f"Product: {product_name}   |   Prepared {date.today():%m/%d/%Y}"
    if author:
        byline += f"   |   {author}"
    _style_run(sub.add_run(byline), 8.5, italic=True)
    _spacing(sub, 0, 4)

    current_cap = None
    shown = 0
    inline_title = None
    inline_links = []
    last_name = ""

    for raw in brief_markdown.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        stripped = line.strip()
        is_bold_heading = stripped.startswith("**") and stripped.endswith("**")

        if is_bold_heading or line.startswith("## "):
            _flush_links(doc, inline_title, inline_links)
            inline_title, inline_links, last_name = None, [], ""

        if is_bold_heading:
            current_cap = SECTION_CAPS.get(stripped)
            shown = 0
            if stripped in INLINE_SECTIONS:
                inline_title = INLINE_SECTIONS[stripped]
                continue

        is_bullet = bool(re.match(r"^\s*[*-]\s+", line))

        if inline_title and is_bullet:
            match = re.match(r"^\s*[*-]\s+\[(.+?)\]\((.+?)\)", line)
            if match:
                label, url = match.groups()
                name = label.split(":")[0] if inline_title == "Competitor pages" else ""
                inline_links.append((_short_label(url, name if name != last_name else ""), url))
                last_name = name
                continue

        if current_cap is not None and is_bullet:
            shown += 1
            if shown > current_cap:
                continue

        if line.startswith("## "):
            heading = doc.add_paragraph(style="Heading 2")
            _style_run(heading.add_run(line[3:].strip()), 11)
            _spacing(heading, 5, 1)
        elif is_bullet:
            text = re.sub(r"^\s*[*-]\s+", "", line)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.first_line_indent = Inches(-0.15)
            _style_run(p.add_run("\u2022  "), BODY_PT)
            _add_inline(p, text)
            _spacing(p, 0, 1)
        else:
            p = doc.add_paragraph()
            _add_inline(p, line.strip())
            _spacing(p, 0, 2)

    _flush_links(doc, inline_title, inline_links)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()