"""Turn a finished account brief into a PowerPoint deck for the sales meeting."""
import io
import math
import re
from datetime import date

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from cap_931.report_generator import _short_label

BLUE = RGBColor(0x03, 0x9B, 0xE5)
DARK = RGBColor(0x26, 0x32, 0x38)
GRAY = RGBColor(0x75, 0x75, 0x75)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

CHARS_PER_LINE = 72  # roughly how many 24 pt characters fit across the body box
MAX_LINES_PER_SLIDE = 10
MAX_NEWS_ON_SLIDE = 6


def _clean(text):
    """Remove markdown marks from a line of text."""
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", text)
    text = text.replace("**", "")
    text = re.sub(r"(?<!\w)\*(?!\s)([^*]+?)\*(?!\w)", r"\1", text)
    return text.strip()


def _split_sections(markdown):
    """Return [(title, [lines])] for each '## ' section of the brief."""
    sections = []
    for raw in markdown.splitlines():
        if raw.startswith("## "):
            sections.append((raw[3:].strip(), []))
        elif sections and raw.strip():
            sections[-1][1].append(raw.rstrip())
    return sections


def _to_bullets(lines):
    """Turn a section's lines into short bullet texts."""
    bullets = []
    for line in lines:
        bullet = re.match(r"^\s*[*-]\s+(.*)$", line)
        if bullet:
            text = _clean(bullet.group(1))
            if text:
                bullets.append(text)
        else:
            for sentence in re.split(r"(?<=[.!?])\s+(?=[A-Z])", line.strip()):
                text = _clean(sentence)
                if text:
                    bullets.append(text)
    return bullets


def _chunk(bullets):
    """Split bullets into groups that fit on one slide."""
    chunks, current, lines = [], [], 0
    for bullet in bullets:
        needed = max(1, math.ceil(len(bullet) / CHARS_PER_LINE))
        if current and lines + needed > MAX_LINES_PER_SLIDE:
            chunks.append(current)
            current, lines = [], 0
        current.append(bullet)
        lines += needed
    if current:
        chunks.append(current)
    return chunks


def _parse_links(lines):
    """Read the Article Links section: {category: [(label, url, suffix)]}."""
    groups, category = {}, None
    for line in lines:
        stripped = line.strip()
        if stripped.startswith("**") and stripped.endswith("**"):
            category = stripped.strip("*")
            groups[category] = []
            continue
        match = re.match(r"^\s*[*-]\s+\[(.+?)\]\((.+?)\)\s*(.*)$", line)
        if match and category:
            groups[category].append(match.groups())
    return groups


def _add_title_bar(prs, slide, text):
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.1))
    band.fill.solid()
    band.fill.fore_color.rgb = BLUE
    band.line.fill.background()
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.15), prs.slide_width - Inches(1.2), Inches(0.8))
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.paragraphs[0].alignment = PP_ALIGN.LEFT
    run = frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(32)
    run.font.bold = True
    run.font.color.rgb = WHITE


def _add_footer(prs, slide, text):
    box = slide.shapes.add_textbox(Inches(0.6), prs.slide_height - Inches(0.55), prs.slide_width - Inches(1.2), Inches(0.35))
    box.text_frame.word_wrap = True
    box.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(12)
    run.font.color.rgb = GRAY


def _make_bullet(paragraph):
    """Give a paragraph a real bullet with a hanging indent."""
    properties = paragraph._p.get_or_add_pPr()
    properties.set("marL", "342900")
    properties.set("indent", "-342900")
    bullet = properties.makeelement(qn("a:buChar"), {"char": "\u2022"})
    properties.append(bullet)


def _body_frame(prs, slide):
    box = slide.shapes.add_textbox(Inches(0.7), Inches(1.5), prs.slide_width - Inches(1.4), Inches(5.2))
    frame = box.text_frame
    frame.word_wrap = True
    return frame


def _add_bullet_slide(prs, title, bullets, footer):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_title_bar(prs, slide, title)
    frame = _body_frame(prs, slide)
    for k, text in enumerate(bullets):
        paragraph = frame.paragraphs[0] if k == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(12)
        run = paragraph.add_run()
        run.text = text
        run.font.size = Pt(24)
        run.font.color.rgb = DARK
        _make_bullet(paragraph)
    _add_footer(prs, slide, footer)


def _add_sources_slide(prs, groups, footer):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_title_bar(prs, slide, "Sources")
    frame = _body_frame(prs, slide)
    first = True

    def new_paragraph():
        nonlocal first
        paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
        first = False
        paragraph.space_after = Pt(8)
        return paragraph

    def add_run(paragraph, text, bold=False, url=None):
        run = paragraph.add_run()
        run.text = text
        run.font.size = Pt(16)
        run.font.bold = bold
        if url:
            run.hyperlink.address = url
        else:
            run.font.color.rgb = DARK

    for category, items in groups.items():
        if not items:
            continue
        if category in ("Company pages", "Competitor pages"):
            paragraph = new_paragraph()
            add_run(paragraph, f"{category}:  ", bold=True)
            last_name = ""
            for k, (label, url, _suffix) in enumerate(items):
                name = label.split(":")[0] if category == "Competitor pages" else ""
                if k:
                    add_run(paragraph, "  |  ")
                add_run(paragraph, _short_label(url, name if name != last_name else ""), url=url)
                last_name = name
        else:
            paragraph = new_paragraph()
            add_run(paragraph, category, bold=True)
            for label, url, suffix in items[:MAX_NEWS_ON_SLIDE]:
                paragraph = new_paragraph()
                add_run(paragraph, label, url=url)
                if suffix:
                    add_run(paragraph, f"  {suffix}")
    _add_footer(prs, slide, footer)


def build_deck(brief_markdown, company_name, product_name, author=""):
    """Return the bytes of a .pptx deck built from the brief."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    footer = f"{company_name}  |  {product_name}"

    # Title slide
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(2.4), prs.slide_width, Inches(2.2))
    band.fill.solid()
    band.fill.fore_color.rgb = BLUE
    band.line.fill.background()
    box = slide.shapes.add_textbox(Inches(0.8), Inches(2.6), prs.slide_width - Inches(1.6), Inches(1.8))
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.paragraphs[0].alignment = PP_ALIGN.LEFT
    run = frame.paragraphs[0].add_run()
    run.text = f"Account Brief: {company_name}"
    run.font.size = Pt(44)
    run.font.bold = True
    run.font.color.rgb = WHITE
    sub = box.text_frame.add_paragraph()
    sub.alignment = PP_ALIGN.LEFT
    run = sub.add_run()
    details = f"{product_name}   |   Prepared {date.today():%m/%d/%Y}"
    if author:
        details += f"   |   {author}"
    run.text = details
    run.font.size = Pt(18)
    run.font.color.rgb = WHITE

    sections = _split_sections(brief_markdown)
    link_groups = {}
    for title, lines in sections:
        if title.lower() == "article links":
            link_groups = _parse_links(lines)
            continue
        bullets = _to_bullets(lines)
        for k, chunk in enumerate(_chunk(bullets)):
            slide_title = title if k == 0 else f"{title} (cont.)"
            _add_bullet_slide(prs, slide_title, chunk, footer)

    if any(link_groups.values()):
        _add_sources_slide(prs, link_groups, footer)

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()