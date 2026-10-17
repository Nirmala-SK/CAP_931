import io

from docx import Document


def read_uploaded_file(uploaded_file, max_chars=8000):
    """Return the plain text of an uploaded .txt, .pdf, .docx or .pptx file."""
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()
    text = ""

    try:
        if name.endswith(".txt"):
            text = data.decode("utf-8", errors="ignore")
        elif name.endswith(".pdf"):
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            text = " ".join((page.extract_text() or "") for page in reader.pages)
        elif name.endswith(".docx"):
            doc = Document(io.BytesIO(data))
            text = " ".join(p.text for p in doc.paragraphs)
        elif name.endswith(".pptx"):
            from pptx import Presentation

            slides = Presentation(io.BytesIO(data)).slides
            parts = []
            for slide in slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        parts.append(shape.text_frame.text)
            text = " ".join(parts)
    except Exception as error:
        print(f"Could not read {uploaded_file.name}: {error}")
        return ""

    return " ".join(text.split())[:max_chars]