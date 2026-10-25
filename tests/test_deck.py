import io

from pptx import Presentation

from cap_931.deck_generator import build_deck

BRIEF = """## Company Strategy
JPMorganChase operates in over 100 markets. It employs about 65,000 technologists. The firm is migrating banking systems to the cloud.

## Leadership Information
*   **Jamie Dimon**: Chairman and Chief Executive Officer
*   **Lori Beer**: Global Chief Information Officer

## Talking Points
*   First point about *security*.
*   Second point.

## Article Links

**Company pages**
- [https://www.jpmorganchase.com/](https://www.jpmorganchase.com/)
- [https://www.jpmorganchase.com/about/leadership](https://www.jpmorganchase.com/about/leadership)

**Recent news**
- [Big headline](https://news.example/1) (Reuters, Tue, 07 Apr 2026)
"""


def slides_text(data):
    prs = Presentation(io.BytesIO(data))
    return [
        " ".join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)
        for slide in prs.slides
    ]


def test_deck_has_title_section_and_sources_slides():
    texts = slides_text(build_deck(BRIEF, "jpmorganchase.com", "SentinelShield", "N"))
    assert "Account Brief: jpmorganchase.com" in texts[0]
    assert any("Company Strategy" in t for t in texts)
    assert any("Talking Points" in t for t in texts)
    assert "Sources" in texts[-1]
    assert not any("Article Links" in t for t in texts)


def test_markdown_marks_are_removed():
    texts = " ".join(slides_text(build_deck(BRIEF, "x.com", "P")))
    assert "**" not in texts
    assert "*security*" not in texts
    assert "Jamie Dimon: Chairman" in texts


def test_long_sections_are_split_across_slides():
    long_brief = "## Leadership Information\n" + "\n".join(
        f"*   **Person {i}**: " + "A long job title " * 8 for i in range(20)
    )
    texts = slides_text(build_deck(long_brief, "x.com", "P"))
    assert sum("Leadership Information" in t for t in texts) >= 2
    assert any("(cont.)" in t for t in texts)


def test_sources_slide_has_clickable_links():
    prs = Presentation(io.BytesIO(build_deck(BRIEF, "x.com", "P")))
    addresses = []
    for shape in prs.slides[-1].shapes:
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.hyperlink.address:
                        addresses.append(run.hyperlink.address)
    assert "https://news.example/1" in addresses
    assert "https://www.jpmorganchase.com/about/leadership" in addresses